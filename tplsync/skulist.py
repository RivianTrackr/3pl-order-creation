"""Read an Allowed SKUs spreadsheet (.xlsx or .csv): client name, SKU and optional description per row."""

import csv
import io
from typing import List, Optional, Tuple

Row = Tuple[str, str, Optional[str]]


class SheetError(ValueError):
    pass


def _columns(header: List[str]) -> Optional[Tuple[int, int, Optional[int]]]:
    """Positions of the client, SKU and description columns, found by their header names."""
    names = [str(h or "").strip().casefold() for h in header]
    client = next((i for i, n in enumerate(names) if "client" in n or "customer" in n), None)
    sku = next((i for i, n in enumerate(names) if "sku" in n), None)
    desc = next((i for i, n in enumerate(names) if "desc" in n), None)
    return (client, sku, desc) if client is not None and sku is not None else None


def _rows(table: List[List[object]]) -> List[Row]:
    table = [r for r in table if any(str(c or "").strip() for c in r)]
    if not table:
        raise SheetError("The file has no rows.")
    cols = _columns([str(c or "") for c in table[0]])
    body = table[1:] if cols else table
    client_i, sku_i, desc_i = cols or (0, 1, 2)
    rows = []
    for r in body:
        cell = lambda i: str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else ""  # noqa: E731
        if cell(client_i) and cell(sku_i):
            rows.append((cell(client_i), cell(sku_i), cell(desc_i) or None))
    if not rows:
        raise SheetError("No rows with both a client name and a SKU were found.")
    return rows


def read_sheet(filename: str, data: bytes) -> List[Row]:
    name = (filename or "").casefold()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return _rows([list(r) for r in csv.reader(io.StringIO(text))])
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # noqa: BLE001 - any unreadable workbook
            raise SheetError(f"The spreadsheet couldn't be read ({type(exc).__name__}).") from exc
        return _rows([list(r) for r in sheet.iter_rows(values_only=True)])
    raise SheetError("Upload an .xlsx or .csv file.")
